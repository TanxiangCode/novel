#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
《消失的灵气》Markdown 稿件 → Word 导出工具

用法：
    python3 tools/export_word.py                 # 导出 4 个分类文档 + 1 个全量合集
    python3 tools/export_word.py --only 正文 设定  # 只导指定分类
    python3 tools/export_word.py --check          # 只跑内容完整性核查（不写文件）

日常维护：新增章节或设定文件后直接重跑即可，文件会被自动归入分组（见「分组规则」段）。
只有想让某个文件排在前面、或换分组时，才需要改下面的 PRIORITY / GROUP_OF 提示表。
"""
import argparse
import datetime
import re
import sys
from pathlib import Path
from xml.etree import ElementTree as ET

import markdown
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / 'novel-project'
OUT = ROOT / 'Word导出'
BOOK = '消失的灵气'
ORDER = ['正文', '大纲', '设定', '状态']

EXTS = ['tables', 'fenced_code', 'sane_lists', 'attr_list']
BLOCKY = ('ul', 'ol', 'table', 'pre', 'blockquote', 'div', 'p')
BREAK = '\x00'

# ===================== 分组规则（按需调整） =====================
# 分类内分组的展示顺序；未列出的分组排在最后
GROUP_ORDER = {
    '设定': ['总纲与规范', '角色', '势力', '修炼', '科技', '地理', '社会', '通用'],
    '大纲': ['全书规划', '分卷大纲与章纲', '场景卡', '专项表与补充'],
    '状态': ['追踪表与台账', '分析与经验沉淀'],
}
# 组内靠前排的文件（按 stem 匹配，其余按自然序）
PRIORITY = {
    '设定': {'角色': ['角色总览', '陈默', '林晚', '苏薇', '叶青鸾', '周远征', '顾承安',
                      '陆衡', '宋晚棠', '沈铁柱', '方岳', '老槐', '代号鹰', '青玄真人'],
             '势力': ['太虚门详设', '天衍宗详设', '万法宗详设', '玄机阁详设', '幽冥殿详设',
                      '三大帝国详设'],
             '修炼': ['功法体系', '阵法体系', '丹药与丹道体系', '法宝与兵器体系',
                      '灵植与灵兽体系', '秘境与遗迹体系'],
             '科技': ['两界物理与科学修仙模型', '科技侧全域监控与情报侦测系统',
                      '蓝星造物与装备名录', '各国计划与装备路线', '方舟基地与玄门学堂'],
             '地理': ['玄洲地理总图', '地名志'],
             '社会': ['历史与纪年', '社会生活与经济', '饮食体系'],
             '总纲与规范': ['世界观', '全文总纲', '设定汇总目录', '设定建设规划',
                            '文风基准', '已固化锚点清单', '势力关系图']},
    '状态': {'追踪表与台账': ['伏笔追踪表', '角色状态追踪', '道具持有表', '数量账登记表',
                              '情绪物触发清单', '行动清单执行台账']},
    '大纲': {'全书规划': ['章纲总纲']},
}

CHAPTER_RE = re.compile(r'第\s*(\d+)\s*章')
VOLUME_RE = re.compile(r'^第(\d+)卷(大纲|章纲)$')


def natkey(path):
    stem = path.stem
    m = CHAPTER_RE.search(stem)
    if m:
        return (0, int(m.group(1)), stem)
    return (1, 0, re.sub(r'\d+', lambda d: '%08d' % int(d.group()), stem))


def group_of(folder, path):
    """决定一个文件归入哪个分组。"""
    rel = path.relative_to(BASE / folder)
    if folder == '正文':
        return '章节正文'
    if folder == '设定':
        return rel.parts[0] if len(rel.parts) > 1 else '总纲与规范'
    if folder == '状态':
        stem = path.stem
        return '分析与经验沉淀' if ('分析报告' in stem or '经验' in stem) else '追踪表与台账'
    stem = path.stem
    if stem.startswith('场景卡'):
        return '场景卡'
    if VOLUME_RE.match(stem):
        return '分卷大纲与章纲'
    if stem in ('章纲总纲',) or '定稿表' in stem or '执行计划' in stem:
        return '全书规划'
    return '专项表与补充'


def volume_key(path):
    m = VOLUME_RE.match(path.stem)
    return (int(m.group(1)), 0 if m.group(2) == '大纲' else 1) if m else (99, 99)


def discover():
    """扫描四个目录，返回 {分类: [(分组标题, [相对路径, ...]), ...]}。"""
    docs = {}
    for folder in ORDER:
        files = sorted((BASE / folder).rglob('*.md'), key=natkey)
        buckets = {}
        for f in files:
            buckets.setdefault(group_of(folder, f), []).append(f)
        groups = []
        for gname in sorted(buckets, key=lambda g: GROUP_ORDER.get(folder, []).index(g)
                            if g in GROUP_ORDER.get(folder, []) else 99):
            pref = PRIORITY.get(folder, {}).get(gname, [])
            items = sorted(buckets[gname],
                           key=lambda p: (pref.index(p.stem) if p.stem in pref else len(pref),
                                          volume_key(p), natkey(p)))
            groups.append((gname, [str(p.relative_to(BASE)) for p in items]))
        if folder == '正文' and groups:
            nums = [int(m.group(1)) for g in groups for f in g[1]
                    for m in [CHAPTER_RE.search(f)] if m]
            if nums:
                groups[0] = ('章节正文（第%d—%d章）' % (min(nums), max(nums)), groups[0][1])
        docs[folder] = groups
    return docs


# ===================== Markdown 结构修复 =====================
DELIM = re.compile(r'^\s*\|?\s*:?-{2,}.*\|\s*$')
SPLIT_PIPE = re.compile(r'(?<!\\)\|')


def cells_of(line):
    s = line.strip()
    if s.startswith('|'):
        s = s[1:]
    if s.endswith('|'):
        s = s[:-1]
    return [c.strip() for c in SPLIT_PIPE.split(s)]


def repair_tables(text):
    """修两处会被 Python-Markdown 静默丢弃的结构问题：
    1) 表格行折到下一行（下一行不以 | 开头）→ 折回上一行的末格，余下部分另起一行；
    2) 某行列数多于表头 → 该行之后的单元格被截断 → 按最大列数补齐表头与分隔行。"""
    lines = text.split('\n')
    out, i, fence = [], 0, False
    while i < len(lines):
        line = lines[i]
        if line.lstrip().startswith('```'):
            fence = not fence
            out.append(line)
            i += 1
            continue
        if fence or not line.lstrip().startswith('|') or i + 1 >= len(lines) \
                or not DELIM.match(lines[i + 1]):
            out.append(line)
            i += 1
            continue
        block, j = [], i
        while j < len(lines) and lines[j].strip() and (
                lines[j].lstrip().startswith('|') or '|' in lines[j]):
            raw = lines[j]
            if not raw.lstrip().startswith('|'):
                head, _, rest = raw.partition('|')
                if not block:
                    break
                last = block.pop()
                head = head.strip()
                block.append(last + ' ' + head if head else last)
                raw = '|' + rest
            block.append(raw)
            j += 1
        header = cells_of(block[0])
        width = max([len(cells_of(r)) for r in block] + [len(header)])
        rebuilt = ['| ' + ' | '.join(header + [''] * (width - len(header))) + ' |',
                   '|' + '---|' * width]
        for row in block[2:]:
            cs = cells_of(row)
            cs += [''] * (width - len(cs))
            rebuilt.append('| ' + ' | '.join(cs) + ' |')
        out.extend(rebuilt)
        i = j
    return '\n'.join(out)


def parse_md(path):
    text = repair_tables(path.read_text(encoding='utf-8'))
    text = re.sub(r'<br\s*/?>', '<br/>', text)
    html = markdown.markdown(text, extensions=EXTS, output_format='xhtml')
    return ET.fromstring('<root>' + html + '</root>')


def top_heading(path):
    for line in path.read_text(encoding='utf-8').splitlines():
        m = re.match(r'^#\s+(.*?)\s*$', line)
        if m:
            return m.group(1)
    return path.stem


# ===================== docx 写入辅助 =====================
def style_fonts(style, latin, east):
    style.font.name = latin
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), latin)
    rf.set(qn('w:eastAsia'), east)


def run_east(run, east):
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:eastAsia'), east)


def shade(pr, fill):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:fill'), fill)
    pr.append(el)


def hrule(paragraph):
    borders = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), '8')
    b.set(qn('w:color'), 'AAAAAA')
    borders.append(b)
    paragraph._p.get_or_add_pPr().append(borders)


def inline(el, fmt=None):
    """展开一个内联节点：自身 text + 子节点（含各自 tail）。"""
    fmt = fmt or {}
    out = []
    if el.text:
        out.append((el.text, fmt))
    for ch in el:
        if ch.tag == 'br':
            out.append((BREAK, fmt))
        else:
            nf = dict(fmt)
            if ch.tag == 'strong':
                nf['bold'] = True
            elif ch.tag in ('em', 'i'):
                nf['italic'] = True
            elif ch.tag == 'code':
                nf['code'] = True
            elif ch.tag in ('del', 's', 'strike'):
                nf['strike'] = True
            out.extend(inline(ch, nf))
        if ch.tail:
            out.append((ch.tail, fmt))
    return out


def mixed(el):
    """容器（li / td / th）：返回 (内联 runs, 块级子节点)，保留 el.text 与每个子节点的 tail。"""
    runs, blocks = [], []
    if el.text:
        runs.append((el.text, {}))
    for ch in el:
        if ch.tag in BLOCKY:
            blocks.append(ch)
        else:
            runs.extend(inline(ch))
        if ch.tail:
            runs.append((ch.tail, {}))
    return runs, blocks


class Builder:
    def __init__(self, prose=False, demote=1):
        self.prose = prose
        self.demote = demote
        self.base = 12.0 if prose else 10.5
        self.doc = Document()
        self.break_before = 0
        self.first = True
        self.files = 0
        self.setup()

    def setup(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.6)
        st = self.doc.styles
        n = st['Normal']
        style_fonts(n, 'Times New Roman', '宋体')
        n.font.size = Pt(self.base)
        n.font.color.rgb = RGBColor(0, 0, 0)
        pf = n.paragraph_format
        pf.line_spacing = 1.6 if self.prose else 1.3
        pf.space_after = Pt(6 if self.prose else 3)
        pf.space_before = Pt(0)
        for lvl, size in {1: 20, 2: 15, 3: 12.5, 4: 11.5, 5: 11, 6: 10.5}.items():
            h = st['Heading %d' % lvl]
            style_fonts(h, 'Times New Roman', '黑体')
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0, 0, 0)
            hp = h.paragraph_format
            hp.space_before = Pt(21 - 2 * lvl)
            hp.space_after = Pt(11 - lvl)
            hp.line_spacing = 1.3
        for name in ('Quote', 'Intense Quote'):
            try:
                q = st[name]
            except KeyError:
                continue
            style_fonts(q, 'Times New Roman', '楷体')
            q.font.size = Pt(10.5)
            q.font.italic = False
            q.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # ---------- 原子输出 ----------
    def add_runs(self, p, runs, size=None, east='宋体'):
        for text, fmt in runs:
            if text == BREAK:
                p.add_run().add_break()
                continue
            pieces = text.split('\n')
            for i, piece in enumerate(pieces):
                if i:
                    p.add_run().add_break()
                if not piece.strip():
                    if piece:
                        p.add_run(piece)
                    continue
                r = p.add_run(piece)
                if fmt.get('bold'):
                    r.bold = True
                if fmt.get('italic'):
                    r.italic = True
                if fmt.get('strike'):
                    r.font.strike = True
                if fmt.get('code'):
                    r.font.name = 'Courier New'
                    r.font.size = Pt((size or self.base) - 1)
                    run_east(r, '宋体')
                else:
                    run_east(r, east)
                    if size:
                        r.font.size = Pt(size)

    def para(self, style=None):
        p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        if self.break_before and not self.first:
            p.paragraph_format.page_break_before = True
        self.break_before = 0
        self.first = False
        return p

    def heading(self, lvl, text_runs):
        self.add_runs(self.para('Heading %d' % max(1, min(lvl, 6))), text_runs, east='黑体')

    # ---------- 块级 ----------
    def block(self, el, quote=False, indent=None):
        tag = el.tag
        if re.fullmatch(r'h[1-6]', tag):
            self.heading(int(tag[1]) + self.demote, inline(el))
        elif tag == 'p':
            p = self.para('Quote' if quote else None)
            if indent is not None:
                p.paragraph_format.left_indent = indent
            if self.prose:
                p.paragraph_format.first_line_indent = Pt(self.base * 2)
            self.add_runs(p, inline(el))
        elif tag == 'blockquote':
            children = list(el)
            if children:
                for sub in children:
                    self.block(sub, quote=True, indent=indent)
            else:
                self.add_runs(self.para('Quote'), inline(el))
        elif tag in ('ul', 'ol'):
            self.list_of(el, 0, tag == 'ol', indent, quote)
        elif tag == 'table':
            self.table(el)
        elif tag == 'pre':
            self.codeblock(el)
        elif tag == 'hr':
            p = self.para()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            hrule(p)
        elif tag in ('root', 'div', 'body', 'dl'):
            for sub in el:
                self.block(sub, quote=quote, indent=indent)
        elif tag in ('dt', 'dd'):
            runs = inline(el)
            if tag == 'dt':
                runs = [(t, dict(f, bold=True)) for t, f in runs]
            p = self.para()
            p.paragraph_format.left_indent = (indent or Cm(0)) + Cm(0.5)
            self.add_runs(p, runs)
        else:
            self.add_runs(self.para(), inline(el))

    def list_of(self, el, depth, ordered, indent, quote):
        base = (indent or Cm(0)) + Cm(0.7 * depth)
        num = int(el.get('start') or 1)
        for li in el.findall('li'):
            if li.get('value') and str(li.get('value')).isdigit():
                num = int(li.get('value'))
            prefix = ('%d. ' % num) if ordered else ('·  ' if depth else '•  ')
            num += 1
            runs, blocks = mixed(li)
            while runs and runs[-1][0] in (BREAK, ''):
                runs.pop()
            p = self.para()
            pf = p.paragraph_format
            pf.left_indent = base + Cm(0.75) + (Cm(0.5) if quote else Cm(0))
            pf.first_line_indent = Cm(-0.75)
            pf.space_after = Pt(2)
            if runs:
                self.add_runs(p, [(prefix, {})])
                self.add_runs(p, runs)
            for sub in blocks:
                if sub.tag in ('ul', 'ol'):
                    self.list_of(sub, depth + 1, sub.tag == 'ol', indent, quote)
                elif sub.tag == 'p':
                    q = self.para()
                    q.paragraph_format.left_indent = base + Cm(1.5)
                    q.paragraph_format.space_after = Pt(2)
                    self.add_runs(q, inline(sub))
                else:
                    self.block(sub, quote=quote, indent=base + Cm(1.5))

    def codeblock(self, el):
        code = el.find('code')
        src = code if code is not None else el
        lines = (src.text or '').split('\n')
        while lines and not lines[-1].strip():
            lines.pop()
        for line in lines:
            p = self.para()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.5)
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.1
            shade(p._p.get_or_add_pPr(), 'F5F5F5')
            r = p.add_run(line if line.strip() else ' ')
            r.font.name = 'Courier New'
            r.font.size = Pt(8.5)
            run_east(r, '宋体')

    def table(self, el):
        rows = []
        for tr in el.iter('tr'):
            cells = []
            for tc in tr:
                runs, blocks = mixed(tc)
                for sub in blocks:
                    runs.extend(self.flat(sub))
                while runs and runs[-1][0] == BREAK:
                    runs.pop()
                cells.append(runs)
            if cells:
                rows.append(cells)
        if not rows:
            return
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.style = 'Table Grid'
        table.autofit = True
        self.break_before = 1
        self.para().paragraph_format.space_after = Pt(0)
        for ri, cells in enumerate(rows):
            for ci in range(ncols):
                p = table.cell(ri, ci).paragraphs[0]
                self.add_runs(p, cells[ci] if ci < len(cells) else [('', {})], size=8.5)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                if ri == 0:
                    shade(table.cell(0, ci)._tc.get_or_add_tcPr(), 'E9E9E9')
                    for r in p.runs:
                        r.bold = True

    @staticmethod
    def flat(el):
        """把表格单元子里的块级内容压成带换行的内联序列。"""
        if el.tag in ('p', 'dt', 'dd'):
            return inline(el) + [(BREAK, {})]
        if el.tag in ('ul', 'ol'):
            out = []
            num = int(el.get('start') or 1)
            for li in el.findall('li'):
                head = ('%d. ' % num) if el.tag == 'ol' else '· '
                num += 1
                runs, blocks = mixed(li)
                out.append((head, {}))
                out.extend(runs)
                for sub in blocks:
                    out.extend(Builder.flat(sub))
                out.append((BREAK, {}))
            return out
        if el.tag == 'pre':
            code = el.find('code')
            txt = (code if code is not None else el).text or ''
            return [(txt.rstrip('\n'), {'code': True}), (BREAK, {})]
        if el.tag == 'blockquote':
            out = []
            for sub in el:
                out.extend(Builder.flat(sub))
            return out
        return inline(el)

    # ---------- 文档级 ----------
    def title_page(self, cn, parts):
        tp = self.para()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tp.paragraph_format.space_before = Pt(96)
        r = tp.add_run('《%s》' % BOOK)
        r.font.size = Pt(30)
        r.bold = True
        run_east(r, '黑体')
        sp = self.para()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = sp.add_run(cn)
        r.font.size = Pt(20)
        run_east(r, '黑体')
        total = sum(len(fs) for _, groups in parts for _, fs in groups)
        mp = self.para()
        mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = mp.add_run('合并 %d 个源文件　·　导出 %s'
                       % (total, datetime.date.today().isoformat()))
        r.font.size = Pt(10.5)
        run_east(r, '宋体')
        r = self.para('Heading 1').add_run('目录')
        run_east(r, '黑体')
        for ptitle, groups in parts:
            if ptitle:
                pp = self.para()
                pp.paragraph_format.space_before = Pt(10)
                r = pp.add_run(ptitle)
                r.bold = True
                r.font.size = Pt(12)
                run_east(r, '黑体')
            for gtitle, files in groups:
                gp = self.para()
                gp.paragraph_format.left_indent = Cm(0.6 if ptitle else 0)
                gp.paragraph_format.space_before = Pt(6)
                r = gp.add_run(gtitle)
                r.bold = True
                r.font.size = Pt(11)
                run_east(r, '黑体')
                for rel in files:
                    ip = self.para()
                    ip.paragraph_format.left_indent = Cm(1.4 if ptitle else 0.8)
                    ip.paragraph_format.space_after = Pt(1)
                    r = ip.add_run(top_heading(BASE / rel))
                    r.font.size = Pt(10.5)
                    run_east(r, '宋体')
        self.first = True
        self.break_before = 1

    def emit(self, label, parts):
        self.title_page(label, parts)
        for ptitle, groups in parts:
            if ptitle:
                self.break_before = 1
                self.heading(1, [(ptitle, {})])
            for gtitle, files in groups:
                self.break_before = 1
                self.heading(2 if ptitle else 1, [(gtitle, {})])
                for rel in files:
                    self.break_before = 1
                    self.files += 1
                    self.block(parse_md(BASE / rel))


def build(label, parts, prose, demote, filename):
    planned = [f for _, groups in parts for _, fs in groups for f in fs]
    if len(planned) != len(set(planned)):
        raise SystemExit('%s 分组内有重复文件' % label)
    tops = {Path(f).parts[0] for f in planned}
    present = sorted(str(p.relative_to(BASE)) for t in tops for p in (BASE / t).rglob('*.md'))
    if set(present) != set(planned):
        raise SystemExit('%s 与磁盘不一致\n仅磁盘有: %s\n仅计划有: %s' % (
            label, sorted(set(present) - set(planned)), sorted(set(planned) - set(present))))
    b = Builder(prose=prose, demote=demote)
    b.emit(label, parts)
    b.doc.core_properties.title = '《%s》%s' % (BOOK, label)
    b.doc.core_properties.comments = '由 novel-project 的 %d 个 markdown 合并导出' % b.files
    out = OUT / filename
    b.doc.save(str(out))
    print('OK  %-30s %3d 文件  %6.0f KB' % (out.name, b.files, out.stat().st_size / 1024.0))
    return out


# ===================== 完整性核查 =====================
HAN = re.compile(r'[一-鿿]')


def han(s):
    return ''.join(HAN.findall(s))


def check(docs, built=None):
    """逐行比对：源文件里的汉字序列必须能在导出的 docx 中连续找到。"""
    failed = False
    targets = {cn: [Path(f).parts[0] for _, fs in docs[cn] for f in fs] for cn in docs}
    if built is None:
        targets['全量合集'] = ORDER
    else:
        targets = {cn: v for cn, v in targets.items() if cn in built}
        if '全量合集' in built:
            targets['全量合集'] = ORDER
    for cn, folders in targets.items():
        path = OUT / ('%s-%s.docx' % (BOOK, cn))
        if not path.exists():
            print('MISSING %s' % path.name)
            failed = True
            continue
        d = Document(str(path))
        doc = han(''.join(el.text for el in d.element.body.iter()
                          if el.tag == qn('w:t') and el.text))
        src_total, misses = 0, []
        for fd in dict.fromkeys(folders):
            for f in sorted((BASE / fd).rglob('*.md')):
                for line in f.read_text(encoding='utf-8').splitlines():
                    key = han(line)
                    src_total += len(key)
                    if len(key) >= 4 and key not in doc:
                        misses.append((f.name, line))
        ratio = len(doc) / src_total
        print('%-6s 汉字 源%-7d docx%-7d 比值%.4f | 未命中 %d 行'
              % (cn, src_total, len(doc), ratio, len(misses)))
        for name, line in misses[:5]:
            print('     · %s: %s' % (name, line[:70]))
        failed = failed or bool(misses)
    print('\n核查通过：内容无缺失' if not failed else '\n核查失败：见上方未命中行')
    return 1 if failed else 0


def main():
    ap = argparse.ArgumentParser(description='导出《%s》Word 文档' % BOOK)
    ap.add_argument('--only', nargs='*', choices=ORDER, help='只导这些分类')
    ap.add_argument('--check', action='store_true', help='只跑完整性核查')
    ap.add_argument('--manifest', action='store_true', help='打印分组结果后退出')
    args = ap.parse_args()

    if not BASE.is_dir():
        raise SystemExit('找不到源目录 %s（脚本需放在仓库 tools/ 下）' % BASE)
    docs = discover()

    if args.manifest:
        for cn in ORDER:
            print('【%s】' % cn)
            for g, fs in docs[cn]:
                print('  %s (%d)' % (g, len(fs)))
                for f in fs:
                    print('      %s' % f)
        return 0
    if args.check:
        return check(docs)

    OUT.mkdir(parents=True, exist_ok=True)
    picked = args.only or ORDER
    built = []
    for cn in picked:
        build(cn, [(None, docs[cn])], prose=(cn == '正文'), demote=1,
              filename='%s-%s.docx' % (BOOK, cn))
        built.append(cn)
    if len(picked) == len(ORDER):
        build('全量合集', [(cn, docs[cn]) for cn in ORDER], prose=False, demote=2,
              filename='%s-全量合集.docx' % BOOK)
        built.append('全量合集')
    return check({cn: docs[cn] for cn in docs if cn in built} if len(built) < len(ORDER)
                 else docs, built)


if __name__ == '__main__':
    sys.exit(main())
